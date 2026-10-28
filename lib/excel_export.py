import json
import re
import sys
from collections import defaultdict
from datetime import date, timedelta
from functools import cmp_to_key
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


root = Path(sys.argv[1]).resolve()
config = json.loads((root / "config.json").read_text(encoding="utf-8"))
history_dir = root / "data" / "history"
monthly_dir = root / "data" / "monthly"
monthly_dir.mkdir(parents=True, exist_ok=True)

cities = ["石家庄", "保定", "邯郸", "沧州", "廊坊", "衡水", "唐山", "雄安"]
city_order = {city: index for index, city in enumerate(cities)}
retention_days = int(config.get("lookbackDays", config.get("retentionDays", 183)))
cutoff = max(
    date.fromisoformat(config.get("projectStartDate", "2026-07-01")),
    date.today() - timedelta(days=retention_days),
).isoformat()


def project_key(title: str) -> str:
    value = re.sub(r"^【[^】]+】", "", title)
    value = re.sub(
        r"[（(]\s*(?:二次|三次|重新招标|重新发布|含招标文件|预公示|暂定项目名称)\s*[)）]",
        "",
        value,
    )
    value = re.sub(r"(?:更正|澄清|变更)(?:公告)?$", "", value)
    value = re.sub(
        r"(?:公开招标公告|招标预公示|招标文件提前公示|预公示|招标计划|"
        r"资格预审|中标候选人公示|中标结果|结果公告|招标公告|"
        r"公开招标|招标|结果)(?:公告)?$",
        "",
        value,
    )
    return re.sub(r"[\s：:，,。.\-—_（）()]", "", value)


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def display_title(value: object) -> str:
    return re.sub(r"^【[^】]+】\s*", "", clean(value))


def design_stage(title: str) -> bool:
    included = re.search(r"设计", title)
    excluded = re.search(
        r"施工招标|土建|施工总承包|工程总承包|EPC|设计施工总承包|采购施工|"
        r"监理招标|监理服务|设备采购|设备招标|采购设备|物资采购|物资招标|"
        r"采购物资|材料采购|材料招标|采购材料|货物采购|设备及材料|"
        r"设计.{0,24}(?:安装|采购).{0,24}设备|"
        r"(?:设备|系统|软件|材料|物资|货物).{0,30}(?:设计|采购|安装)|"
        r"采购及安装|设备安装|机电安装|安装工程|"
        r"中标候选人|中标结果|结果公告|成交公告|终止公告|废标公告|流标公告",
        title,
        re.IGNORECASE,
    )
    return bool(included and not excluded)


records = {}
for path in history_dir.glob("*.json"):
    try:
        summary = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        continue
    for group in summary.get("tenders", {}).get("cityGroups", []):
        city = group.get("city", "")
        for record in [*group.get("qualified", []), *group.get("pending", [])]:
            if clean(record.get("publicationDate", ""))[:10] < cutoff:
                continue
            if (
                record.get("recordType") != "approval"
                and not design_stage(clean(record.get("title", "")))
            ):
                continue
            item = {**record, "city": city}
            key = project_key(item.get("title", "")) or item.get("sourceUrl", "")
            existing = records.get(key)
            if existing is None or item.get("publicationDate", "") >= existing.get(
                "publicationDate", ""
            ):
                records[key] = item

by_month = defaultdict(list)
for record in records.values():
    month = clean(record.get("publicationDate", ""))[:7]
    if re.fullmatch(r"\d{4}-\d{2}", month):
        by_month[month].append(record)

header = [
    "所在城市",
    "项目名称",
    "项目阶段",
    "业主单位",
    "项目概况",
    "近期进展情况",
    "投资金额",
    "计划招标时间",
    "信息来源链接",
    "核验状态",
]
widths = [12, 52, 14, 30, 62, 28, 16, 30, 58, 12]
monthly_index = []

for month, rows in sorted(by_month.items(), reverse=True):
    def compare(left, right):
        left_city = city_order.get(left.get("city", ""), 99)
        right_city = city_order.get(right.get("city", ""), 99)
        if left_city != right_city:
            return left_city - right_city
        left_stage = 0 if left.get("recordType") == "approval" else 1
        right_stage = 0 if right.get("recordType") == "approval" else 1
        if left_stage != right_stage:
            return left_stage - right_stage
        if left.get("publicationDate", "") != right.get("publicationDate", ""):
            return -1 if left.get("publicationDate", "") > right.get("publicationDate", "") else 1
        return (right.get("investmentAmountYuan") or 0) - (
            left.get("investmentAmountYuan") or 0
        )

    rows.sort(key=cmp_to_key(compare))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = month
    sheet.append(header)
    for record in rows:
        sheet.append(
            [
                record.get("city", ""),
                display_title(record.get("title", "")),
                "立项" if record.get("recordType") == "approval" else "招标设计",
                record.get("owner", ""),
                record.get("overview", ""),
                record.get("recentProgress", ""),
                record.get("investmentAmount", ""),
                record.get("plannedTenderTime", ""),
                record.get("sourceUrl", ""),
                (
                    "符合条件"
                    if record.get("qualified")
                    else "新立项·待核实"
                    if record.get("recordType") == "approval"
                    else "待核实"
                ),
            ]
        )
        source_cell = sheet.cell(sheet.max_row, 9)
        source_cell.hyperlink = record.get("sourceUrl", "")
        source_cell.style = "Hyperlink"

    header_fill = PatternFill("solid", fgColor="0E6B4E")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 90
    sheet.row_dimensions[1].height = 28

    info = workbook.create_sheet("说明")
    info.append(["私人秘书 · 河北工程项目月报"])
    info.append(["统计月份", month])
    info.append(["项目数量", len(rows)])
    info.append(
        [
            "公告阶段",
            "仅保留2026年7月1日以来、有效期半年内的初步设计、方案设计、勘察设计或设计招标；"
            "排除施工、土建、EPC、监理、设备/材料/物资/货物采购、结果及终止公告。",
        ]
    )
    info.append(["筛选范围", "石家庄、保定、邯郸、沧州、廊坊、衡水、唐山、雄安"])
    info.append(["专业范围", "公路、铁路、市政、房建、水利、电力、矿山"])
    info.append(["金额规则", "严格大于1.5亿元；未披露金额标记待核实"])
    info.append(["资金规则", "只采信官方原文明确的资金落实描述"])
    info.append(
        [
            "新立项来源",
            "河北省发展和改革委员会行政许可公示；新立项项目单列并标记待核实。",
        ]
    )
    info.append(
        [
            "设计公告来源",
            "河北省招标投标服务平台（招标公告、招标计划、招标文件提前公示）、"
            "河北省公共资源交易中心、雄安新区公共资源交易服务平台；跨平台重复公告合并。",
        ]
    )
    info.append(["说明", "待核实项目不代表不符合金额条件，需结合招标文件附件核验。"])
    info["A1"].font = Font(size=15, bold=True, color="0E6B4E")
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 78
    for row in info.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = monthly_dir / f"{month}.xlsx"
    workbook.save(output)
    monthly_index.append(
        {
            "month": month,
            "label": f"{month[:4]}年{int(month[5:])}月",
            "file": f"data/monthly/{month}.xlsx",
            "count": len(rows),
        }
    )

(root / "data" / "months.json").write_text(
    json.dumps(monthly_index, ensure_ascii=False, indent=2) + "\n",
    encoding="utf-8",
)
for stale_file in monthly_dir.glob("*.xlsx"):
    if stale_file.stem not in by_month:
        stale_file.unlink()
print(json.dumps(monthly_index, ensure_ascii=False))
